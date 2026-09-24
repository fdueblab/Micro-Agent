"""文件处理服务：上传保存、ZIP 解压、URL 下载、临时文件清理。"""

from __future__ import annotations

import base64
import os
import re
import shutil
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Optional

import httpx
from fastapi import HTTPException, UploadFile
from loguru import logger


# ── Filename sanitisation ────────────────────────────────────────────────
# 多端点（/code_analysis、/service_packaging、/aml_auto_generate 等）会把
# multipart 上传的 filename 头直接拼到工作区路径里。攻击者可以构造
# filename = '../../etc/cron.d/x' 之类的路径元素来逃逸 dest_dir：当工作区
# 里恰好存在同名首段目录时（长时间运行的服务里很常见），写入会落到工作区
# 外部；在 Windows 上 ``\\`` 是路径分隔符，攻击直接成功；即便不能逃逸，
# 未清洗的文件名也会污染工作区并破坏后续 cleanup 逻辑。统一在保存前提取
# basename 并替换危险字符。
_UNSAFE_FILENAME_CHARS = re.compile(r"[^A-Za-z0-9._-]")


def _safe_filename(name: Optional[str], default: str = "upload") -> str:
    """从用户提供的文件名中提取一个安全 basename。

    - 去掉路径分隔符（``/`` 与 ``\\``），只保留 ``os.path.basename``/``Path.name``
    - 把 ``..``、空字符串、全为点号的名字（``.``、``...`` 等）替换成默认值
    - 把除 ``A-Za-z0-9._-`` 之外的字符替换成 ``_``，避免 shell/命令注入隐患
    - 限制总长度，防止过长文件名 DoS 或破坏文件系统
    """
    raw = (name or "").strip()
    if not raw:
        return default

    # 同时切掉 Windows 风格分隔符——Path.name 在 POSIX 上不会拆 ``\\``。
    raw = raw.replace("\\", "/")
    base = os.path.basename(raw)

    # ``..``、``.`` 等仅由点号组成的名字视为非法。
    if not base or set(base) <= {"."}:
        return default

    cleaned = _UNSAFE_FILENAME_CHARS.sub("_", base)

    # 再次防御：替换后仍可能退化成 ``.``/``..``。
    if not cleaned or set(cleaned) <= {"."}:
        return default

    # 文件系统多数对单个组件限制 255 字节，预留 ts/前缀空间。
    return cleaned[:200]


async def save_upload(upload: UploadFile, dest_dir: Path) -> Path:
    dest_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    safe_name = _safe_filename(upload.filename)
    dest = dest_dir / f"{ts}_{safe_name}"

    # 兜底校验：确保最终路径仍在 dest_dir 内，防止未来逻辑改动引入回归。
    resolved_dest = dest.resolve()
    resolved_root = dest_dir.resolve()
    if not resolved_dest.is_relative_to(resolved_root):
        raise HTTPException(400, "非法的上传文件名")

    content = await upload.read()
    dest.write_bytes(content)
    logger.info(f"文件已保存: {dest} ({len(content)} bytes)")
    return dest


def extract_zip(zip_path: Path, dest_dir: Path) -> Path:
    dest_dir.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path, "r") as zf:
        zf.extractall(dest_dir)
    logger.info(f"ZIP 已解压: {zip_path} -> {dest_dir}")
    return dest_dir


async def download_from_url(url: str, dest_dir: Path) -> Path:
    """从 URL 下载文件到 dest_dir，返回保存路径。"""
    from urllib.parse import urlparse

    dest_dir.mkdir(parents=True, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    parsed = urlparse(url)
    raw_name = os.path.basename(parsed.path) or f"download_{ts}.zip"
    # 同样清洗远程响应里推断出的文件名，避免「URL path = ../../x」之类的攻击
    # 经由同一个 dest_dir 拼接落到工作区外（与 save_upload 风险对称）。
    filename = _safe_filename(raw_name, default=f"download_{ts}.zip")
    dest = dest_dir / f"{ts}_{filename}"

    async with httpx.AsyncClient(timeout=120) as client:
        resp = await client.get(url)
        if resp.status_code != 200:
            raise HTTPException(400, f"无法从 URL 下载文件，状态码: {resp.status_code}")
        dest.write_bytes(resp.content)

    logger.info(f"文件已下载: {url} -> {dest}")
    return dest


def resolve_project_dir(saved: Path, workspace: Path) -> str:
    """处理上传的 ZIP/PY 文件，返回项目目录路径。"""
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    ext = saved.suffix.lower()

    if ext == ".zip":
        extract_dir = workspace / f"{ts}_extracted"
        extract_zip(saved, extract_dir)
        items = [
            i for i in os.listdir(extract_dir)
            if not i.startswith(".") and i not in {"__MACOSX", "Thumbs.db", "desktop.ini"}
        ]
        if len(items) == 1 and (extract_dir / items[0]).is_dir():
            return str(extract_dir / items[0])
        return str(extract_dir)

    if ext == ".py":
        project_dir = workspace / f"{ts}_extracted"
        project_dir.mkdir(parents=True, exist_ok=True)
        shutil.copy2(str(saved), str(project_dir / saved.name))
        return str(project_dir)

    raise HTTPException(400, f"不支持的文件类型: {ext}，仅支持 .zip 和 .py")


def find_main_file(directory: str) -> str:
    """在目录中查找主 Python 文件。"""
    priority = ("main.py", "app.py", "server.py", "run.py", "start.py", "__main__.py")
    for name in priority:
        if os.path.isfile(os.path.join(directory, name)):
            return name
    for f in os.listdir(directory):
        if f.endswith(".py") and os.path.isfile(os.path.join(directory, f)):
            return f
    return ""


async def resolve_file_or_url(
    file: Optional[UploadFile],
    file_url: Optional[str],
    dest_dir: Path,
) -> Path:
    """统一处理 '文件上传 OR URL 下载' 二选一逻辑，返回本地文件路径。"""
    has_file = file is not None and getattr(file, "filename", None)
    has_url = file_url is not None and file_url.strip() != ""

    if not has_file and not has_url:
        raise HTTPException(400, "必须提供文件上传或文件 URL")

    if has_file:
        return await save_upload(file, dest_dir)
    return await download_from_url(file_url.strip(), dest_dir)


def pack_directory_as_zip_base64(directory: str) -> dict:
    """将目录打包为 ZIP 并返回 base64 编码结果。"""
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    zip_path = Path(directory).parent / f"{ts}_service_package.zip"
    folder_name = os.path.basename(directory)

    try:
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
            for root, _dirs, files in os.walk(directory):
                for f in files:
                    fp = os.path.join(root, f)
                    arcname = os.path.relpath(fp, os.path.dirname(directory))
                    zf.write(fp, arcname)

        with open(zip_path, "rb") as f:
            content = base64.b64encode(f.read()).decode("utf-8")

        return {
            "filename": zip_path.name,
            "content": content,
            "type": "zip",
        }
    finally:
        if zip_path.exists():
            zip_path.unlink()


def read_paper_content(file_path: str) -> str:
    """从 PDF / DOC / DOCX 文件中提取文本内容。"""
    _, ext = os.path.splitext(file_path.lower())
    content = ""
    try:
        if ext == ".pdf":
            from PyPDF2 import PdfReader
            reader = PdfReader(file_path)
            for page in reader.pages:
                content += (page.extract_text() or "")
        elif ext in (".doc", ".docx"):
            from docx import Document
            doc = Document(file_path)
            for para in doc.paragraphs:
                content += para.text + "\n"
    except Exception as e:
        logger.warning(f"提取文件内容失败 ({file_path}): {e}")
    return content


def read_reference_text(file_path: str, max_chars: int = 4000) -> str:
    """读取参考资料文件文本内容。

    支持 PDF / DOC / DOCX（复用 read_paper_content）、纯文本(.txt/.md)、
    代码文件(.py/.ipynb)。返回截断到 max_chars 的文本。
    """
    _, ext = os.path.splitext(file_path.lower())
    content = ""
    try:
        if ext in (".pdf", ".doc", ".docx"):
            content = read_paper_content(file_path)
        elif ext in (".txt", ".md", ".py", ".ipynb", ".json", ".csv"):
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
        elif ext == ".zip":
            # 解压后拼接其中的代码/文本文件内容
            parts = []
            try:
                with zipfile.ZipFile(file_path, "r") as zf:
                    for name in zf.namelist():
                        if name.endswith((".py", ".txt", ".md", ".json")) and not name.startswith("__MACOSX"):
                            try:
                                parts.append(f"# {name}\n" + zf.read(name).decode("utf-8", errors="replace"))
                            except Exception:
                                continue
                        if sum(len(p) for p in parts) > max_chars:
                            break
            except Exception as e:
                logger.warning(f"读取 ZIP 参考资料失败 ({file_path}): {e}")
            content = "\n\n".join(parts)
        else:
            logger.info(f"参考资料类型暂不支持文本提取，跳过: {ext}")
    except Exception as e:
        logger.warning(f"提取参考资料内容失败 ({file_path}): {e}")
    return (content or "")[:max_chars]


async def fetch_url_text(url: str, max_chars: int = 4000) -> str:
    """抓取 URL 的可读文本（简单去标签）。失败返回空字符串。"""
    import re

    try:
        async with httpx.AsyncClient(timeout=30, follow_redirects=True) as client:
            resp = await client.get(url, headers={"User-Agent": "Mozilla/5.0 Micro-Agent"})
            if resp.status_code != 200:
                logger.warning(f"抓取 URL 失败 [{resp.status_code}]: {url}")
                return ""
            text = resp.text
    except Exception as e:
        logger.warning(f"抓取 URL 异常 ({url}): {e}")
        return ""

    # 粗略去除 script/style 和 HTML 标签
    text = re.sub(r"(?is)<(script|style)[^>]*>.*?</\1>", " ", text)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text[:max_chars]


def cleanup_paths(*paths: str | Path) -> None:
    """安全清理临时文件和目录。"""
    for p in paths:
        try:
            p = Path(p)
            if not p.exists():
                continue
            if p.is_dir():
                shutil.rmtree(p)
            else:
                p.unlink()
            logger.debug(f"已清理: {p}")
        except Exception as e:
            logger.warning(f"清理 {p} 失败: {e}")


def parse_dataset_file(file_path: str, max_rows: int = 20) -> dict:
    """解析数据集文件，提取元信息（格式、列名、样例行、标签分布等）。

    支持 CSV / Excel / JSON / TXT / PDF 格式。
    返回 dict 包含 format, columns, total_rows, sample_rows, label_distribution, raw_text。
    """
    _, ext = os.path.splitext(file_path.lower())
    result: dict = {"format": ext.lstrip("."), "file_name": os.path.basename(file_path)}

    try:
        if ext == ".csv":
            import csv
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                rows = list(reader)
            if not rows:
                result["raw_text"] = "(空文件)"
                return result
            result["columns"] = rows[0]
            result["total_rows"] = len(rows) - 1
            result["sample_rows"] = rows[1 : max_rows + 1]
            result["raw_text"] = _rows_to_text(rows[0], rows[1 : max_rows + 1])
            result["label_distribution"] = _detect_label_distribution(rows)

        elif ext in (".xlsx", ".xls"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
                ws = wb.active
                rows = [[str(cell) if cell is not None else "" for cell in row] for row in ws.iter_rows(values_only=True)]
                wb.close()
            except Exception:
                result["raw_text"] = "(无法解析 Excel 文件，请改用 CSV 格式)"
                return result
            if not rows:
                result["raw_text"] = "(空文件)"
                return result
            result["columns"] = rows[0]
            result["total_rows"] = len(rows) - 1
            result["sample_rows"] = rows[1 : max_rows + 1]
            result["raw_text"] = _rows_to_text(rows[0], rows[1 : max_rows + 1])
            result["label_distribution"] = _detect_label_distribution(rows)

        elif ext == ".json":
            import json as json_mod
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                data = json_mod.load(f)
            if isinstance(data, list):
                result["total_rows"] = len(data)
                result["sample_rows"] = data[:max_rows]
                if data and isinstance(data[0], dict):
                    result["columns"] = list(data[0].keys())
                result["raw_text"] = json_mod.dumps(data[:max_rows], ensure_ascii=False, indent=2)
            else:
                result["raw_text"] = json_mod.dumps(data, ensure_ascii=False, indent=2)[:3000]

        elif ext == ".txt":
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
            lines = content.strip().split("\n")
            result["total_rows"] = len(lines)
            result["raw_text"] = "\n".join(lines[:max_rows])

        elif ext == ".pdf":
            content = read_paper_content(file_path)
            if content:
                lines = [l.strip() for l in content.split("\n") if l.strip()]
                result["total_rows"] = len(lines)
                result["raw_text"] = "\n".join(lines[:80])
            else:
                result["raw_text"] = "(PDF 内容提取为空)"

        else:
            result["raw_text"] = f"(不支持的数据集格式: {ext})"

    except Exception as e:
        logger.warning(f"解析数据集文件失败 ({file_path}): {e}")
        result["raw_text"] = f"(解析失败: {e})"

    return result


def _rows_to_text(columns: list, rows: list) -> str:
    """将表格数据转为可读文本。"""
    header = " | ".join(str(c) for c in columns)
    lines = [header, "-" * len(header)]
    for row in rows:
        lines.append(" | ".join(str(c) for c in row))
    return "\n".join(lines)


def _detect_label_distribution(rows: list) -> dict:
    """尝试从表格数据中检测标签列及其分布。

    启发式：寻找名字含 type/label/class/category 的列，统计各值计数。
    """
    if len(rows) < 2:
        return {}
    header = [str(c).lower().strip() for c in rows[0]]
    label_keywords = ("type", "label", "class", "category", "action", "tag")
    label_col_idx = None
    for i, col in enumerate(header):
        if any(kw in col for kw in label_keywords):
            label_col_idx = i
            break
    if label_col_idx is None:
        return {}
    distribution: dict[str, int] = {}
    for row in rows[1:]:
        if label_col_idx < len(row):
            val = str(row[label_col_idx]).strip()
            if val:
                distribution[val] = distribution.get(val, 0) + 1
    return distribution
